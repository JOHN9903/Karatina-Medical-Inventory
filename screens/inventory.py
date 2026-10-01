import sys
import os
from datetime import datetime

# ---------------------------------------------------------
# PROJECT ROOT
# ---------------------------------------------------------
PROJECT_ROOT = os.path.dirname(
    os.path.dirname(os.path.abspath(__file__))
)

if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)


# ---------------------------------------------------------
# PYQT
# ---------------------------------------------------------
from PySide6.QtWidgets import (
    QWidget,
    QDialog,
    QVBoxLayout,
    QHBoxLayout,
    QFormLayout,
    QLabel,
    QLineEdit,
    QComboBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QMessageBox,
    QHeaderView,
    QGroupBox,
    QFileDialog,
    QAbstractItemView,
)

from PySide6.QtCore import Qt


# ---------------------------------------------------------
# DATABASE
# ---------------------------------------------------------
from database.connection import get_connection


# ---------------------------------------------------------
# OPENPYXL
# ---------------------------------------------------------
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


# =========================================================
# ADD EQUIPMENT DIALOG
# =========================================================
class AddEquipmentDialog(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setWindowTitle("Add Medical Equipment")
        self.setMinimumWidth(500)

        self.setup_ui()
        self.load_equipment_names()
        self.load_departments()

    # -----------------------------------------------------
    # UI
    # -----------------------------------------------------
    def setup_ui(self):

        layout = QVBoxLayout(self)

        title = QLabel("Add Medical Equipment")
        title.setStyleSheet("""
            QLabel {
                font-size: 20px;
                font-weight: bold;
                color: #0B5D4B;
                padding: 8px;
            }
        """)

        layout.addWidget(title)

        form = QFormLayout()

        self.equipment_name_combo = QComboBox()
        self.equipment_name_combo.setMinimumHeight(35)

        self.model_input = QLineEdit()
        self.model_input.setPlaceholderText("Enter model number")

        self.serial_input = QLineEdit()
        self.serial_input.setPlaceholderText("Enter serial number")

        self.manufacturer_input = QLineEdit()
        self.manufacturer_input.setPlaceholderText("Enter manufacturer")

        self.year_input = QLineEdit()
        self.year_input.setPlaceholderText("e.g. 2024")

        self.department_combo = QComboBox()
        self.department_combo.setMinimumHeight(35)

        self.status_combo = QComboBox()
        self.status_combo.addItems([
            "Functional",
            "Non-Functional",
            "Under Maintenance",
            "Decommissioned"
        ])

        self.remarks_input = QLineEdit()
        self.remarks_input.setPlaceholderText("Optional remarks")

        form.addRow("Equipment Name:", self.equipment_name_combo)
        form.addRow("Model Number:", self.model_input)
        form.addRow("Serial Number:", self.serial_input)
        form.addRow("Manufacturer:", self.manufacturer_input)
        form.addRow("Year of Manufacture:", self.year_input)
        form.addRow("Department:", self.department_combo)
        form.addRow("Status:", self.status_combo)
        form.addRow("Remarks:", self.remarks_input)

        layout.addLayout(form)

        buttons = QHBoxLayout()

        self.save_button = QPushButton("Save Equipment")
        self.cancel_button = QPushButton("Cancel")

        self.save_button.setStyleSheet("""
            QPushButton {
                background-color: #0B5D4B;
                color: white;
                font-weight: bold;
                padding: 10px 18px;
                border-radius: 5px;
            }

            QPushButton:hover {
                background-color: #084638;
            }
        """)

        self.cancel_button.setStyleSheet("""
            QPushButton {
                padding: 10px 18px;
                border-radius: 5px;
            }
        """)

        self.save_button.clicked.connect(self.save_equipment)
        self.cancel_button.clicked.connect(self.reject)

        buttons.addStretch()
        buttons.addWidget(self.save_button)
        buttons.addWidget(self.cancel_button)

        layout.addLayout(buttons)

    # -----------------------------------------------------
    # LOAD EQUIPMENT NAMES
    # -----------------------------------------------------
    def load_equipment_names(self):

        try:
            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT equipment_name_id, equipment_name
                FROM equipment_names
                WHERE is_active = 1
                ORDER BY equipment_name
            """)

            rows = cursor.fetchall()

            self.equipment_name_combo.clear()

            for row in rows:
                self.equipment_name_combo.addItem(
                    row[1],
                    row[0]
                )

            cursor.close()
            connection.close()

        except Exception as e:
            QMessageBox.critical(
                self,
                "Database Error",
                f"Could not load equipment names.\n\n{e}"
            )

    # -----------------------------------------------------
    # LOAD DEPARTMENTS
    # -----------------------------------------------------
    def load_departments(self):

        try:
            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT department_id, department_name
                FROM departments
                WHERE is_active = 1
                ORDER BY department_name
            """)

            rows = cursor.fetchall()

            self.department_combo.clear()

            for row in rows:
                self.department_combo.addItem(
                    row[1],
                    row[0]
                )

            cursor.close()
            connection.close()

        except Exception as e:
            QMessageBox.critical(
                self,
                "Database Error",
                f"Could not load departments.\n\n{e}"
            )

    # -----------------------------------------------------
    # SAVE EQUIPMENT
    # -----------------------------------------------------
    def save_equipment(self):

        equipment_name_id = self.equipment_name_combo.currentData()
        department_id = self.department_combo.currentData()

        model_number = self.model_input.text().strip()
        serial_number = self.serial_input.text().strip()
        manufacturer = self.manufacturer_input.text().strip()
        year_text = self.year_input.text().strip()
        status = self.status_combo.currentText()
        remarks = self.remarks_input.text().strip()

        if not equipment_name_id:
            QMessageBox.warning(
                self,
                "Missing Information",
                "Please select an equipment name."
            )
            return

        if not department_id:
            QMessageBox.warning(
                self,
                "Missing Information",
                "Please select a department."
            )
            return

        year_of_manufacture = None

        if year_text:

            try:
                year_of_manufacture = int(year_text)

                if year_of_manufacture < 1900 or year_of_manufacture > datetime.now().year:
                    raise ValueError

            except ValueError:
                QMessageBox.warning(
                    self,
                    "Invalid Year",
                    "Please enter a valid year of manufacture."
                )
                return

        try:

            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                INSERT INTO equipment_inventory
                (
                    equipment_name_id,
                    model_number,
                    serial_number,
                    manufacturer,
                    year_of_manufacture,
                    department_id,
                    status,
                    date_added,
                    remarks,
                    is_active
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    CURDATE(),
                    %s,
                    1
                )
            """, (
                equipment_name_id,
                model_number or None,
                serial_number or None,
                manufacturer or None,
                year_of_manufacture,
                department_id,
                status,
                remarks or None
            ))

            connection.commit()

            cursor.close()
            connection.close()

            QMessageBox.information(
                self,
                "Success",
                "Equipment added successfully."
            )

            self.accept()

        except Exception as e:

            try:
                connection.rollback()
            except Exception:
                pass

            QMessageBox.critical(
                self,
                "Database Error",
                f"Could not save equipment.\n\n{e}"
            )


# =========================================================
# INVENTORY WINDOW
# =========================================================
class EquipmentInventoryWindow(QWidget):

    def __init__(
        self,
        parent=None,
        user_id=1,
        role="Viewer"
    ):

        super().__init__(parent)

        self.user_id = user_id
        self.role = role

        self.setWindowTitle(
            "Medical Equipment Inventory"
        )

        self.resize(1400, 800)

        self.setup_ui()
        self.load_departments()
        self.load_inventory()

    # =====================================================
    # UI
    # =====================================================
    def setup_ui(self):

        main_layout = QVBoxLayout(self)

        # -------------------------------------------------
        # TITLE
        # -------------------------------------------------
        title_layout = QVBoxLayout()

        title = QLabel(
            "Medical Equipment Inventory"
        )

        title.setStyleSheet("""
            QLabel {
                font-size: 26px;
                font-weight: bold;
                color: #0B5D4B;
            }
        """)

        subtitle = QLabel(
            "Karatina Sub-County Hospital — Equipment Inventory Management"
        )

        subtitle.setStyleSheet("""
            QLabel {
                color: #666666;
                font-size: 13px;
            }
        """)

        title_layout.addWidget(title)
        title_layout.addWidget(subtitle)

        main_layout.addLayout(title_layout)

        # -------------------------------------------------
        # FILTERS
        # -------------------------------------------------
        filter_group = QGroupBox("Inventory Filters")

        filter_layout = QHBoxLayout()

        department_label = QLabel("Department:")

        self.department_filter = QComboBox()
        self.department_filter.setMinimumWidth(220)
        self.department_filter.addItem(
            "All Departments",
            None
        )

        status_label = QLabel("Status:")

        self.status_filter = QComboBox()
        self.status_filter.setMinimumWidth(180)

        self.status_filter.addItem(
            "All Statuses",
            None
        )

        self.status_filter.addItems([
            "Functional",
            "Non-Functional",
            "Under Maintenance",
            "Decommissioned"
        ])

        search_label = QLabel("Search:")

        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText(
            "Search equipment, model, serial number, manufacturer..."
        )
        self.search_input.setMinimumWidth(300)

        self.refresh_button = QPushButton("Refresh")

        self.refresh_button.setStyleSheet("""
            QPushButton {
                padding: 8px 15px;
            }
        """)

        filter_layout.addWidget(department_label)
        filter_layout.addWidget(self.department_filter)

        filter_layout.addWidget(status_label)
        filter_layout.addWidget(self.status_filter)

        filter_layout.addWidget(search_label)
        filter_layout.addWidget(self.search_input)

        filter_layout.addWidget(self.refresh_button)

        filter_group.setLayout(filter_layout)

        main_layout.addWidget(filter_group)

        # -------------------------------------------------
        # ACTION BUTTONS
        # -------------------------------------------------
        action_layout = QHBoxLayout()

        self.add_button = QPushButton(
            "+ Add Equipment"
        )

        self.add_button.setStyleSheet("""
            QPushButton {
                background-color: #0B5D4B;
                color: white;
                font-weight: bold;
                padding: 10px 18px;
                border-radius: 5px;
            }

            QPushButton:hover {
                background-color: #084638;
            }
        """)

        self.deactivate_button = QPushButton(
            "Deactivate Selected"
        )

        self.deactivate_button.setStyleSheet("""
            QPushButton {
                background-color: #B91C1C;
                color: white;
                font-weight: bold;
                padding: 10px 18px;
                border-radius: 5px;
            }

            QPushButton:hover {
                background-color: #991B1B;
            }
        """)

        self.export_button = QPushButton(
            "Export to Excel"
        )

        self.export_button.setStyleSheet("""
            QPushButton {
                background-color: #2563EB;
                color: white;
                font-weight: bold;
                padding: 10px 18px;
                border-radius: 5px;
            }

            QPushButton:hover {
                background-color: #1D4ED8;
            }
        """)

        action_layout.addWidget(self.add_button)
        action_layout.addWidget(self.deactivate_button)

        action_layout.addStretch()

        # Export only for Admin and Authorized User
        if self.role in ["Admin", "Authorized User"]:
            action_layout.addWidget(self.export_button)

        main_layout.addLayout(action_layout)

        # -------------------------------------------------
        # TABLE
        # -------------------------------------------------
        self.table = QTableWidget()

        self.table.setColumnCount(10)

        self.table.setHorizontalHeaderLabels([
            "ID",
            "Equipment Name",
            "Category",
            "Model Number",
            "Serial Number",
            "Manufacturer",
            "Year",
            "Department",
            "Status",
            "Remarks"
        ])

        self.table.setSelectionBehavior(
            QAbstractItemView.SelectRows
        )

        self.table.setSelectionMode(
            QAbstractItemView.SingleSelection
        )

        self.table.setEditTriggers(
            QAbstractItemView.NoEditTriggers
        )

        self.table.setAlternatingRowColors(True)

        self.table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeToContents
        )

        self.table.horizontalHeader().setStretchLastSection(True)

        main_layout.addWidget(self.table)

        # -------------------------------------------------
        # CONNECTIONS
        # -------------------------------------------------
        self.department_filter.currentIndexChanged.connect(
            self.load_inventory
        )

        self.status_filter.currentIndexChanged.connect(
            self.load_inventory
        )

        self.search_input.textChanged.connect(
            self.load_inventory
        )

        self.refresh_button.clicked.connect(
            self.load_inventory
        )

        self.add_button.clicked.connect(
            self.add_equipment
        )

        self.deactivate_button.clicked.connect(
            self.deactivate_equipment
        )

        if self.role in ["Admin", "Authorized User"]:
            self.export_button.clicked.connect(
                self.export_to_excel
            )

        # -------------------------------------------------
        # ROLE PERMISSIONS
        # -------------------------------------------------
        if self.role == "Viewer":

            self.add_button.setEnabled(False)
            self.deactivate_button.setEnabled(False)

        elif self.role == "Authorized User":

            self.add_button.setEnabled(True)

            # Authorized users cannot deactivate equipment
            self.deactivate_button.setEnabled(False)

        elif self.role == "Admin":

            self.add_button.setEnabled(True)
            self.deactivate_button.setEnabled(True)

    # =====================================================
    # LOAD DEPARTMENTS
    # =====================================================
    def load_departments(self):

        try:

            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT department_id, department_name
                FROM departments
                WHERE is_active = 1
                ORDER BY department_name
            """)

            rows = cursor.fetchall()

            current_data = self.department_filter.currentData()

            self.department_filter.blockSignals(True)

            self.department_filter.clear()

            self.department_filter.addItem(
                "All Departments",
                None
            )

            selected_index = 0

            for row in rows:

                self.department_filter.addItem(
                    row[1],
                    row[0]
                )

                if row[0] == current_data:
                    selected_index = (
                        self.department_filter.count() - 1
                    )

            self.department_filter.setCurrentIndex(
                selected_index
            )

            self.department_filter.blockSignals(False)

            cursor.close()
            connection.close()

        except Exception as e:

            self.department_filter.blockSignals(False)

            QMessageBox.critical(
                self,
                "Database Error",
                f"Could not load departments.\n\n{e}"
            )

    # =====================================================
    # LOAD INVENTORY
    # =====================================================
    def load_inventory(self):

        try:

            connection = get_connection()
            cursor = connection.cursor()

            query = """
                SELECT
                    ei.equipment_id,
                    en.equipment_name,
                    COALESCE(ec.category_name, 'Uncategorized') AS category_name,
                    ei.model_number,
                    ei.serial_number,
                    ei.manufacturer,
                    ei.year_of_manufacture,
                    d.department_name,
                    ei.status,
                    ei.remarks
                FROM equipment_inventory ei

                INNER JOIN equipment_names en
                    ON ei.equipment_name_id = en.equipment_name_id

                LEFT JOIN equipment_categories ec
                    ON en.category_id = ec.category_id

                INNER JOIN departments d
                    ON ei.department_id = d.department_id

                WHERE ei.is_active = 1
            """

            params = []

            # ---------------------------------------------
            # DEPARTMENT FILTER
            # ---------------------------------------------
            department_id = (
                self.department_filter.currentData()
            )

            if department_id:

                query += """
                    AND ei.department_id = %s
                """

                params.append(department_id)

            # ---------------------------------------------
            # STATUS FILTER
            # ---------------------------------------------
            status = self.status_filter.currentData()

            if status:

                query += """
                    AND ei.status = %s
                """

                params.append(status)

            # ---------------------------------------------
            # SEARCH
            # ---------------------------------------------
            search_text = (
                self.search_input.text().strip()
            )

            if search_text:

                query += """
                    AND (
                        en.equipment_name LIKE %s
                        OR ec.category_name LIKE %s
                        OR ei.model_number LIKE %s
                        OR ei.serial_number LIKE %s
                        OR ei.manufacturer LIKE %s
                        OR d.department_name LIKE %s
                        OR ei.status LIKE %s
                        OR ei.remarks LIKE %s
                    )
                """

                search_value = f"%{search_text}%"

                params.extend([
                    search_value,
                    search_value,
                    search_value,
                    search_value,
                    search_value,
                    search_value,
                    search_value,
                    search_value
                ])

            query += """
                ORDER BY
                    d.department_name,
                    en.equipment_name
            """

            cursor.execute(
                query,
                params
            )

            rows = cursor.fetchall()

            self.table.setRowCount(0)

            for row_data in rows:

                row_number = self.table.rowCount()

                self.table.insertRow(row_number)

                for column_number, value in enumerate(row_data):

                    if value is None:
                        value = ""

                    item = QTableWidgetItem(
                        str(value)
                    )

                    item.setTextAlignment(
                        Qt.AlignLeft | Qt.AlignVCenter
                    )

                    self.table.setItem(
                        row_number,
                        column_number,
                        item
                    )

            cursor.close()
            connection.close()

        except Exception as e:

            QMessageBox.critical(
                self,
                "Database Error",
                f"Could not load inventory.\n\n{e}"
            )

    # =====================================================
    # ADD EQUIPMENT
    # =====================================================
    def add_equipment(self):

        if self.role == "Viewer":

            QMessageBox.warning(
                self,
                "Access Restricted",
                "Viewer accounts cannot add equipment."
            )

            return

        dialog = AddEquipmentDialog(self)

        if dialog.exec():

            self.load_inventory()

    # =====================================================
    # DEACTIVATE EQUIPMENT
    # =====================================================
    def deactivate_equipment(self):

        if self.role != "Admin":

            QMessageBox.warning(
                self,
                "Access Restricted",
                "Only an Administrator can deactivate equipment."
            )

            return

        selected_row = (
            self.table.currentRow()
        )

        if selected_row < 0:

            QMessageBox.warning(
                self,
                "No Selection",
                "Please select an equipment item first."
            )

            return

        equipment_id_item = (
            self.table.item(
                selected_row,
                0
            )
        )

        equipment_name_item = (
            self.table.item(
                selected_row,
                1
            )
        )

        equipment_id = (
            equipment_id_item.text()
        )

        equipment_name = (
            equipment_name_item.text()
        )

        confirmation = QMessageBox.question(
            self,
            "Confirm Deactivation",
            f"Are you sure you want to deactivate:\n\n"
            f"{equipment_name}\n\n"
            f"Equipment ID: {equipment_id}\n\n"
            f"This will remove it from the active inventory.",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if confirmation != QMessageBox.Yes:
            return

        try:

            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                UPDATE equipment_inventory
                SET
                    is_active = 0,
                    date_updated = CURRENT_TIMESTAMP
                WHERE equipment_id = %s
            """, (
                equipment_id,
            ))

            connection.commit()

            cursor.close()
            connection.close()

            QMessageBox.information(
                self,
                "Success",
                "Equipment was deactivated successfully."
            )

            self.load_inventory()

        except Exception as e:

            try:
                connection.rollback()
            except Exception:
                pass

            QMessageBox.critical(
                self,
                "Database Error",
                f"Could not deactivate equipment.\n\n{e}"
            )

    # =====================================================
    # EXPORT TO EXCEL
    # =====================================================
    def export_to_excel(self):

        if self.role not in [
            "Admin",
            "Authorized User"
        ]:

            QMessageBox.warning(
                self,
                "Access Restricted",
                "You do not have permission to export inventory."
            )

            return

        if Workbook is None:

            QMessageBox.critical(
                self,
                "Missing Package",
                "Excel export requires openpyxl.\n\n"
                "Run:\n"
                "pip install openpyxl"
            )

            return

        # -------------------------------------------------
        # GET CURRENT FILTERS
        # -------------------------------------------------
        department_id = (
            self.department_filter.currentData()
        )

        department_name = (
            self.department_filter.currentText()
        )

        status = (
            self.status_filter.currentData()
        )

        status_name = (
            self.status_filter.currentText()
        )

        search_text = (
            self.search_input.text().strip()
        )

        # -------------------------------------------------
        # GET VISIBLE TABLE DATA
        # -------------------------------------------------
        row_count = self.table.rowCount()

        if row_count == 0:

            QMessageBox.information(
                self,
                "Nothing to Export",
                "There are no inventory records matching "
                "the current filters."
            )

            return

        # -------------------------------------------------
        # DEFAULT FILE NAME
        # -------------------------------------------------
        date_string = datetime.now().strftime(
            "%Y-%m-%d"
        )

        safe_department = (
            department_name
            .replace("/", "-")
            .replace("\\", "-")
            .replace(" ", "_")
        )

        safe_status = (
            status_name
            .replace("/", "-")
            .replace("\\", "-")
            .replace(" ", "_")
        )

        if department_id:
            default_name = (
                f"Karatina_Inventory_"
                f"{safe_department}_"
                f"{date_string}.xlsx"
            )

        elif status:
            default_name = (
                f"Karatina_Inventory_"
                f"{safe_status}_"
                f"{date_string}.xlsx"
            )

        else:
            default_name = (
                f"Karatina_Inventory_"
                f"{date_string}.xlsx"
            )

        # -------------------------------------------------
        # SELECT SAVE LOCATION
        # -------------------------------------------------
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Inventory to Excel",
            default_name,
            "Excel Workbook (*.xlsx)"
        )

        if not file_path:
            return

        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:

            # -------------------------------------------------
            # CREATE WORKBOOK
            # -------------------------------------------------
            workbook = Workbook()

            worksheet = workbook.active

            worksheet.title = "Inventory"

            # -------------------------------------------------
            # HOSPITAL HEADER
            # -------------------------------------------------
            worksheet["A1"] = (
                "KARATINA SUB-COUNTY HOSPITAL"
            )

            worksheet["A2"] = (
                "MEDICAL EQUIPMENT INVENTORY REPORT"
            )

            worksheet["A3"] = (
                f"Exported By: {self.role}"
            )

            worksheet["A4"] = (
                "Export Date: "
                + datetime.now().strftime(
                    "%d/%m/%Y %H:%M"
                )
            )

            worksheet["A5"] = (
                f"Department Filter: "
                f"{department_name}"
            )

            worksheet["A6"] = (
                f"Status Filter: "
                f"{status_name}"
            )

            worksheet["A7"] = (
                f"Search Filter: "
                f"{search_text if search_text else 'None'}"
            )

            # -------------------------------------------------
            # HEADER STYLE
            # -------------------------------------------------
            worksheet["A1"].font = Font(
                bold=True,
                size=16
            )

            worksheet["A2"].font = Font(
                bold=True,
                size=13
            )

            for cell in [
                worksheet["A3"],
                worksheet["A4"],
                worksheet["A5"],
                worksheet["A6"],
                worksheet["A7"]
            ]:

                cell.font = Font(
                    size=10
                )

            # -------------------------------------------------
            # TABLE HEADER
            # -------------------------------------------------
            header_row = 9

            headers = [
                "ID",
                "Equipment Name",
                "Category",
                "Model Number",
                "Serial Number",
                "Manufacturer",
                "Year",
                "Department",
                "Status",
                "Remarks"
            ]

            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )

            for column, header in enumerate(
                headers,
                start=1
            ):

                cell = worksheet.cell(
                    row=header_row,
                    column=column,
                    value=header
                )

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                cell.border = thin_border

            # -------------------------------------------------
            # EXPORT TABLE DATA
            # -------------------------------------------------
            for row in range(row_count):

                excel_row = (
                    header_row + 1 + row
                )

                for column in range(10):

                    table_item = (
                        self.table.item(
                            row,
                            column
                        )
                    )

                    value = ""

                    if table_item:
                        value = table_item.text()

                    cell = worksheet.cell(
                        row=excel_row,
                        column=column + 1,
                        value=value
                    )

                    cell.border = thin_border

                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True
                    )

            # -------------------------------------------------
            # COLUMN WIDTHS
            # -------------------------------------------------
            widths = [
                8,
                28,
                20,
                18,
                20,
                22,
                12,
                25,
                20,
                35
            ]

            for index, width in enumerate(
                widths,
                start=1
            ):

                worksheet.column_dimensions[
                    get_column_letter(index)
                ].width = width

            # -------------------------------------------------
            # FREEZE HEADER
            # -------------------------------------------------
            worksheet.freeze_panes = "A10"

            # -------------------------------------------------
            # AUTO FILTER
            # -------------------------------------------------
            worksheet.auto_filter.ref = (
                f"A{header_row}:"
                f"J{header_row + row_count}"
            )

            # -------------------------------------------------
            # SAVE
            # -------------------------------------------------
            workbook.save(
                file_path
            )

            QMessageBox.information(
                self,
                "Export Successful",
                f"Inventory exported successfully.\n\n"
                f"Records exported: {row_count}\n\n"
                f"File:\n{file_path}"
            )

        except Exception as e:

            QMessageBox.critical(
                self,
                "Export Error",
                f"Could not export inventory to Excel.\n\n"
                f"{e}"
            )


# =========================================================
# TEST
# =========================================================
if __name__ == "__main__":

    from PySide6.QtWidgets import QApplication

    app = QApplication(sys.argv)

    window = EquipmentInventoryWindow(
        user_id=1,
        role="Admin"
    )

    window.show()

    sys.exit(app.exec())